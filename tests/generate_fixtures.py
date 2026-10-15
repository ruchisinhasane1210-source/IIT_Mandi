"""Generate minimal test fixtures: 3-page PDF, 3-slide PPTX, and 30-second audio file."""
from __future__ import annotations

import io
import math
import struct
import wave
from pathlib import Path

FIXTURES_DIR = Path(__file__).resolve().parent / "fixtures"
FIXTURES_DIR.mkdir(parents=True, exist_ok=True)


def create_pdf_fixture(output_path: Path) -> Path:
    """Generate a 3-page PDF with text, a table, and a figure drawing."""
    import fitz

    doc = fitz.open()

    # Page 1: Intro Text
    p1 = doc.new_page()
    p1.insert_text(
        fitz.Point(50, 72),
        "Chapter 8: Binary Heaps and Priority Queues\n\n"
        "A binary heap is a complete binary tree stored implicitly in an array.\n"
        "Index i has children at 2i+1 and 2i+2 and parent at floor((i-1)/2).\n"
        "The heap property requires each parent to be smaller than its children for a min-heap.",
        fontsize=12,
    )

    # Page 2: Table
    p2 = doc.new_page()
    p2.insert_text(
        fitz.Point(50, 72),
        "Operation Complexity Table:\n"
        "Operation | Complexity\n"
        "Insert    | O(log n)\n"
        "Extract   | O(log n)\n"
        "Build-Heap| O(n)\n",
        fontsize=12,
    )
    # Draw a small rectangle diagram
    p2.draw_rect(fitz.Rect(50, 200, 250, 300), color=(0, 0, 1), fill=(0.9, 0.9, 1))
    p2.insert_text(fitz.Point(60, 240), "[Diagram: Min-Heap Tree Structure]", fontsize=10)

    # Page 3: Build-heap analysis
    p3 = doc.new_page()
    p3.insert_text(
        fitz.Point(50, 72),
        "Build-Heap Analysis:\n\n"
        "Building a heap from an arbitrary array takes O(n), not O(n log n).\n"
        "Because most nodes sit near the leaves and sift down only a short distance;\n"
        "summing costs telescopes to a linear total.",
        fontsize=12,
    )

    doc.save(output_path)
    doc.close()
    return output_path


def create_pptx_fixture(output_path: Path) -> Path:
    """Generate a 3-slide PPTX with title, body, speaker notes, and shapes."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()

    # Slide 1: Title & Overview
    slide1 = prs.slides.add_slide(prs.slide_layouts[0])
    slide1.shapes.title.text = "Lecture 8: Binary Heaps"
    slide1.placeholders[1].text = "Array-backed complete binary trees and heap operations."
    slide1.notes_slide.notes_text_frame.text = "Speaker note: Remind students about zero-based indexing."

    # Slide 2: Operations
    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    slide2.shapes.title.text = "Heap Property and Sift-Down"
    slide2.placeholders[1].text = "Parent at floor((i-1)/2), children at 2i+1 and 2i+2.\nHeapify restores the invariant."
    slide2.notes_slide.notes_text_frame.text = "Speaker note: Walk through diagram on slide 18."

    # Slide 3: Build-Heap Linear Time
    slide3 = prs.slides.add_slide(prs.slide_layouts[1])
    slide3.shapes.title.text = "Linear Time Heap Construction"
    slide3.placeholders[1].text = "Bottom-up heapify achieves O(n) construction time."
    slide3.notes_slide.notes_text_frame.text = "Speaker note: Emphasize geometric series convergence."

    prs.save(str(output_path))
    return output_path


def create_audio_fixture(output_path: Path) -> Path:
    """Generate a 30-second sine wave audio WAV file."""
    sample_rate = 8000
    duration_s = 30
    freq = 440.0
    num_samples = sample_rate * duration_s

    with wave.open(str(output_path), "wb") as wf:
        wf.setnchannels(1)
        wf.setsampwidth(2)
        wf.setframerate(sample_rate)
        frames = bytearray()
        for i in range(num_samples):
            val = int(32767.0 * 0.2 * math.sin(2.0 * math.pi * freq * (i / sample_rate)))
            frames.extend(struct.pack("<h", val))
        wf.writeframes(frames)

    return output_path


def generate_all_fixtures() -> dict[str, Path]:
    pdf_p = FIXTURES_DIR / "sample_3page.pdf"
    pptx_p = FIXTURES_DIR / "sample_3slide.pptx"
    audio_p = FIXTURES_DIR / "sample_30s.wav"

    create_pdf_fixture(pdf_p)
    create_pptx_fixture(pptx_p)
    create_audio_fixture(audio_p)

    return {
        "pdf": pdf_p,
        "pptx": pptx_p,
        "audio": audio_p,
    }


if __name__ == "__main__":
    paths = generate_all_fixtures()
    print("Generated fixtures:")
    for k, p in paths.items():
        print(f"  {k}: {p} (size={p.stat().st_size} bytes)")
