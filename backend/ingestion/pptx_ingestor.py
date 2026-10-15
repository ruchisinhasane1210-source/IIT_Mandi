"""PPTX Ingestor — Slide-level extraction with speaker notes, tables, and figures.

Uses python-pptx to extract slide content (title + body + speaker notes + tables)
as coherent slide units, plus embedded diagrams/images as captioned figure units.
"""
from __future__ import annotations

import io
from pathlib import Path
from typing import Any

from tracelearn.vision import caption_and_ocr


def extract_pptx(
    path: Path,
    source_id: str,
    *,
    course_id: str = "dsa_6006",
    storage_dir: Path | None = None,
) -> tuple[list[dict], list[dict]]:
    """Extract slide units and figures from a PPTX file.

    Returns:
        (units, figures)
    """
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError:
        return [], []

    units: list[dict] = []
    figures: list[dict] = []

    if storage_dir is None:
        storage_dir = path.parent / "figures"
    storage_dir.mkdir(parents=True, exist_ok=True)

    try:
        prs = Presentation(str(path))
    except Exception:
        return [], []

    for slide_idx, slide in enumerate(prs.slides, start=1):
        title_text = ""
        body_parts: list[str] = []
        table_parts: list[str] = []
        slide_images: list[bytes] = []

        # 1. Inspect shapes on slide
        for shape in slide.shapes:
            # Title
            if shape == slide.shapes.title and shape.has_text_frame:
                title_text = shape.text_frame.text.strip()
            # General text
            elif shape.has_text_frame:
                txt = shape.text_frame.text.strip()
                if txt:
                    body_parts.append(txt)
            # Table
            elif shape.has_table:
                table_grid: list[list[str]] = []
                for row in shape.table.rows:
                    row_txt = [c.text.strip().replace("\n", " ") for c in row.cells]
                    table_grid.append(row_txt)
                if table_grid and table_grid[0]:
                    header = table_grid[0]
                    sep = ["---"] * len(header)
                    tlines = [
                        "| " + " | ".join(header) + " |",
                        "| " + " | ".join(sep) + " |",
                    ]
                    for r in table_grid[1:]:
                        padded = r + [""] * (len(header) - len(r))
                        tlines.append("| " + " | ".join(padded[:len(header)]) + " |")
                    table_parts.append("\n".join(tlines))
            # Image shape
            elif getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE or hasattr(shape, "image"):
                try:
                    img_blob = shape.image.blob
                    # Filter tiny images
                    if len(img_blob) > 2048:
                        slide_images.append(img_blob)
                except Exception:
                    pass

        # 2. Speaker notes
        notes_text = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                notes_text = f"\n(Speaker Notes) {notes}"

        # 3. Assemble slide unit
        merged_sections: list[str] = []
        if title_text:
            merged_sections.append(f"Slide {slide_idx}: {title_text}")
        else:
            merged_sections.append(f"Slide {slide_idx}")

        if body_parts:
            merged_sections.append("\n".join(body_parts))
        if table_parts:
            merged_sections.append("\n".join(table_parts))
        if notes_text:
            merged_sections.append(notes_text)

        slide_content = "\n\n".join(merged_sections).strip()
        deep_link = f"/viewer?source={source_id}&slide={slide_idx}"

        if slide_content:
            units.append({
                "content": slide_content,
                "modality": "slide",
                "unit_type": "slide",
                "slide_number": slide_idx,
                "slide_no": slide_idx,
                "token_count": len(slide_content.split()),
                "deep_link": deep_link,
            })

        # 4. Process images as figures linked to the slide
        for img_i, img_blob in enumerate(slide_images, start=1):
            fig_id = f"fig_{source_id}_s{slide_idx}_{img_i}"
            img_path = storage_dir / f"{fig_id}.png"
            try:
                img_path.write_bytes(img_blob)
            except Exception:
                pass

            caption, ocr_txt = caption_and_ocr(
                img_blob,
                context=slide_content[:200],
                source_type="slide",
                page_or_slide=slide_idx,
            )

            fig_record = {
                "figure_id": fig_id,
                "source_id": source_id,
                "course_id": course_id,
                "page_or_slide_or_ts": f"slide {slide_idx}",
                "image_path": str(img_path),
                "caption": caption,
                "ocr_text": ocr_txt,
                "unit_id": f"u_{fig_id}",
            }
            figures.append(fig_record)

            combined_fig_content = f"Figure on slide {slide_idx}: {caption}"
            if ocr_txt:
                combined_fig_content += f"\nVisible text: {ocr_txt}"

            units.append({
                "id": f"u_{fig_id}",
                "content": combined_fig_content,
                "modality": "figure",
                "unit_type": "figure",
                "slide_number": slide_idx,
                "slide_no": slide_idx,
                "figure_id": fig_id,
                "figure_reference": fig_id,
                "token_count": len(combined_fig_content.split()),
                "deep_link": deep_link,
            })

    return units, figures
